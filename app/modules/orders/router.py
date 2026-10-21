import uuid
from typing import List, Optional
from fastapi import APIRouter, Depends, Query
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from pydantic import BaseModel, ConfigDict
from datetime import datetime

from app.core.database import get_db
from app.modules.orders.models import Order
from app.modules.integrations.kaspi.models import KaspiShop

router = APIRouter(prefix="/orders", tags=["Orders & Analytics"])


class OrderResponse(BaseModel):
    id: uuid.UUID
    external_order_id: str
    status: str
    customer_name: Optional[str] = None
    total_sale_amount: float
    total_cost_amount: float
    kaspi_commission_amount: float
    net_profit: float
    margin_percent: float
    created_at: datetime
    model_config = ConfigDict(from_attributes=True)


class AnalyticsSummary(BaseModel):
    total_orders: int
    gross_revenue: float
    total_cost_of_goods: float
    total_kaspi_commission: float
    net_profit: float
    average_margin_percent: float


@router.get("", response_model=List[OrderResponse])
async def list_orders(
    shop_id: Optional[uuid.UUID] = Query(None),
    status: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db)
):
    stmt = select(Order).order_by(Order.created_at.desc())
    if shop_id:
        stmt = stmt.where(Order.kaspi_shop_id == shop_id)
    if status:
        stmt = stmt.where(Order.status == status)
    
    orders = (await db.execute(stmt)).scalars().all()
    return orders


@router.get("/analytics/summary", response_model=AnalyticsSummary)
async def get_analytics_summary(
    shop_id: Optional[uuid.UUID] = Query(None),
    db: AsyncSession = Depends(get_db)
):
    """
    Returns unit economics and profitability dashboard metrics.
    """
    stmt = select(
        func.count(Order.id).label("total_orders"),
        func.coalesce(func.sum(Order.total_sale_amount), 0).label("gross_revenue"),
        func.coalesce(func.sum(Order.total_cost_amount), 0).label("total_cogs"),
        func.coalesce(func.sum(Order.kaspi_commission_amount), 0).label("total_commission"),
        func.coalesce(func.sum(Order.net_profit), 0).label("net_profit")
    )
    if shop_id:
        stmt = stmt.where(Order.kaspi_shop_id == shop_id)

    res = (await db.execute(stmt)).one()
    
    total_orders = res.total_orders or 0
    revenue = float(res.gross_revenue)
    cogs = float(res.total_cogs)
    commission = float(res.total_commission)
    profit = float(res.net_profit)
    
    avg_margin = (profit / revenue * 100) if revenue > 0 else 0.0

    return AnalyticsSummary(
        total_orders=total_orders,
        gross_revenue=revenue,
        total_cost_of_goods=cogs,
        total_kaspi_commission=commission,
        net_profit=profit,
        average_margin_percent=round(avg_margin, 2)
    )


@router.get("/export/excel")
async def export_orders_to_excel(
    shop_id: Optional[uuid.UUID] = Query(None),
    db: AsyncSession = Depends(get_db)
):
    """
    Exports orders and unit economics report as an Excel (.xlsx) spreadsheet.
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from fastapi.responses import Response

    stmt = select(Order).order_by(Order.created_at.desc())
    if shop_id:
        stmt = stmt.where(Order.kaspi_shop_id == shop_id)
    orders = (await db.execute(stmt)).scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Реестр заказов Kaspi"

    # Header styling
    headers = [
        "№ Заказа Kaspi", "Дата", "Статус", "Покупатель", "Телефон",
        "Сумма продажи (₸)", "Себестоимость (₸)", "Комиссия Kaspi (₸)",
        "Доставка (₸)", "Чистая прибыль (₸)", "Рентабельность (%)"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for ord in orders:
        date_str = ord.created_at.strftime("%d.%m.%Y %H:%M") if ord.created_at else ""
        row = [
            ord.external_order_id,
            date_str,
            ord.status,
            ord.customer_name or "—",
            ord.customer_phone or "—",
            float(ord.total_sale_amount),
            float(ord.total_cost_amount),
            float(ord.kaspi_commission_amount),
            float(ord.delivery_cost),
            float(ord.net_profit),
            float(ord.margin_percent)
        ]
        ws.append(row)
        current_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx in [6, 7, 8, 9, 10]:
                c.number_format = '#,##0 "₸"'
                c.alignment = right_align
            elif col_idx == 11:
                c.number_format = '0.0"%"'
                c.alignment = right_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"kaspi_orders_report_{uuid.uuid4().hex[:8]}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
